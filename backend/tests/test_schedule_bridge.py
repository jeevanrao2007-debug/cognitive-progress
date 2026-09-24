"""Tests for Phase 4: Verified Actuals -> Schedule Export Bridge.

Covers:
1. CSV schedule import with realistic headers
2. XLSX schedule import with realistic headers
3. Malformed schedule handling
4. Verified actual update
5. Planner-review exclusion from export
6. Conflict exclusion from export
7. Rejected exclusion from export
8. Unchanged activity handling
9. Export field correctness (CSV)
10. Export field correctness (XLSX)
11. Export audit and provenance traceability
12. End-to-end verified schedule flow
13. Dependency conflict cannot enter verified export
14. Ambiguous match cannot enter verified export
"""

import csv
import io
import json
import uuid
from datetime import UTC, date, datetime
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.database.models import (
    Activity,
    ActivityDependency,
    ActivityStatus,
    ConflictSeverity,
    ConflictType,
    DependencyType,
    Evidence,
    EvidenceConflict,
    EvidenceInterpretation,
    ImportStatus,
    ObservationMatchCandidate,
    ObservationMatchOutcome,
    ObservationMatchResult,
    PlannerReview,
    Project,
    Reconciliation,
    ReconciliationDecision,
    ReviewDecision,
    ScheduleUpdate,
    SourceType,
)
from app.main import app
from app.modules.ingestion.service import IngestionService
from app.modules.reconciliation.service import ReconciliationService
from app.modules.schedule.bridge_service import ScheduleBridgeService
from app.modules.schedule.schemas import ScheduleDecisionState
from app.modules.schedule.updates import ScheduleUpdateService


@pytest.fixture
def bridge_project(session) -> Project:
    project = Project(name="Refinery Schedule Bridge Project")
    session.add(project)
    session.commit()
    return project


# ---------------------------------------------------------------------------
# 1. Realistic CSV Schedule Import
# ---------------------------------------------------------------------------

def test_realistic_csv_schedule_import(session, bridge_project, tmp_path):
    csv_file = tmp_path / "realistic_schedule.csv"
    csv_file.write_text(
        "activity_id,activity_name,wbs,discipline,planned_start,planned_end,actual_start,actual_end,status,percent_complete,predecessors\n"
        "PIP001,Erect Line 24-XX,1.1,Piping,2026-08-25,2026-09-20,,,not_started,0,\n"
        "HT001,Hydrotest Line 24-XX,1.1,Piping,2026-09-21,2026-09-25,,,not_started,0,PIP001\n"
        "ELEC001,Transformer Foundation,1.2,Electrical,2026-08-20,2026-09-05,2026-08-20,,in_progress,60,\n",
        encoding="utf-8",
    )

    source = IngestionService(session).import_schedule(
        bridge_project.project_id, csv_file, "realistic_schedule.csv", "v1"
    )

    assert source.status == ImportStatus.COMPLETED
    assert source.records_imported == 3

    activities = list(
        session.scalars(
            select(Activity)
            .where(Activity.project_id == bridge_project.project_id)
            .order_by(Activity.external_activity_id)
        )
    )
    assert len(activities) == 3

    elec = next(a for a in activities if a.external_activity_id == "ELEC001")
    assert elec.discipline == "Electrical"
    assert elec.actual_start == date(2026, 8, 20)
    assert elec.actual_progress == 60.0
    assert elec.status == ActivityStatus.IN_PROGRESS

    ht = next(a for a in activities if a.external_activity_id == "HT001")
    assert len(ht.predecessors) == 1
    assert ht.predecessors[0].predecessor.external_activity_id == "PIP001"


# ---------------------------------------------------------------------------
# 2. Realistic XLSX Schedule Import
# ---------------------------------------------------------------------------

def test_realistic_xlsx_schedule_import(session, bridge_project, tmp_path):
    xlsx_file = tmp_path / "schedule_extract.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws.append([
        "Activity ID",
        "Activity Name",
        "WBS",
        "Discipline",
        "Planned Start",
        "Planned End",
        "Actual Start",
        "Actual End",
        "Status",
        "Percent Complete",
        "Predecessors",
    ])
    ws.append([
        "CIV001",
        "Foundation Excavation",
        "1.1",
        "Civil",
        "2026-08-01",
        "2026-08-20",
        "2026-08-01",
        "2026-08-19",
        "completed",
        "100",
        "",
    ])
    ws.append([
        "CIV002",
        "Concrete Pour Unit 1",
        "1.1",
        "Civil",
        "2026-08-21",
        "2026-09-10",
        "",
        "",
        "not_started",
        "0",
        "CIV001",
    ])
    wb.save(xlsx_file)

    source = IngestionService(session).import_schedule(
        bridge_project.project_id, xlsx_file, "schedule_extract.xlsx", "v1"
    )

    assert source.status == ImportStatus.COMPLETED
    assert source.records_imported == 2

    civ1 = session.scalar(
        select(Activity).where(
            Activity.project_id == bridge_project.project_id,
            Activity.external_activity_id == "CIV001",
        )
    )
    assert civ1.status == ActivityStatus.COMPLETED
    assert civ1.actual_start == date(2026, 8, 1)
    assert civ1.actual_end == date(2026, 8, 19)
    assert civ1.actual_progress == 100.0

    civ2 = session.scalar(
        select(Activity).where(
            Activity.project_id == bridge_project.project_id,
            Activity.external_activity_id == "CIV002",
        )
    )
    assert len(civ2.predecessors) == 1
    assert civ2.predecessors[0].predecessor.external_activity_id == "CIV001"


# ---------------------------------------------------------------------------
# 3. Malformed Schedule Handling
# ---------------------------------------------------------------------------

def test_malformed_schedule_handling(session, bridge_project, tmp_path):
    bad_csv = tmp_path / "bad_dates.csv"
    bad_csv.write_text(
        "activity_id,activity_name,discipline,planned_start,planned_end\n"
        "ACT001,Test Bad Date,Civil,not-a-date,2026-09-20\n",
        encoding="utf-8",
    )

    source = IngestionService(session).import_schedule(
        bridge_project.project_id, bad_csv, "bad_dates.csv", "v1"
    )
    assert source.status == ImportStatus.FAILED
    assert source.validation_errors is not None
    assert any("planned_start" in str(e) for e in source.validation_errors)


# ---------------------------------------------------------------------------
# 4. Verified Actual Update via AUTO_ACCEPT
# ---------------------------------------------------------------------------

def test_verified_actual_update(session, bridge_project):
    act = Activity(
        project_id=bridge_project.project_id,
        external_activity_id="PIP001",
        description="Erect Line 24-XX",
        discipline="Piping",
        level=6,
        status=ActivityStatus.NOT_STARTED,
        planned_start=date(2026, 8, 25),
        planned_finish=date(2026, 9, 20),
    )
    session.add(act)
    session.flush()

    rec = Reconciliation(
        activity_id=act.activity_id,
        reconciled_status=ActivityStatus.COMPLETED,
        reconciled_progress=100.0,
        actual_start=date(2026, 8, 25),
        actual_end=date(2026, 9, 12),
        confidence=0.98,
        explanation="Field crew completed erection on Sep 12.",
        decision=ReconciliationDecision.AUTO_ACCEPT,
        supporting_evidence=[{
            "observation_id": str(uuid.uuid4()),
            "source_name": "DPR-2026-09-12.txt",
            "source_type": "daily_report",
        }],
    )
    session.add(rec)
    session.flush()

    # Apply deterministic update
    ScheduleUpdateService(session).apply(
        act, rec, decision_source="automatic:reconciliation", reason=rec.explanation
    )
    session.commit()

    service = ScheduleBridgeService(session)
    preview = service.get_change_preview(bridge_project.project_id)

    assert preview.summary.verified_updates == 1
    assert preview.summary.unchanged == 0

    end_change = next(c for c in preview.changes if c.field == "Actual End")
    assert end_change.decision == ScheduleDecisionState.VERIFIED
    assert end_change.old_value == "—"
    assert end_change.new_value == "2026-09-12"
    assert "DPR-2026-09-12.txt" in end_change.source_evidence


# ---------------------------------------------------------------------------
# 5. Planner-Review Exclusion from Verified Export
# ---------------------------------------------------------------------------

def test_planner_review_exclusion_from_verified_export(session, bridge_project):
    act = Activity(
        project_id=bridge_project.project_id,
        external_activity_id="PIP002",
        description="Hydrotest Line 18-AA",
        discipline="Piping",
        level=6,
        status=ActivityStatus.NOT_STARTED,
    )
    session.add(act)
    session.flush()

    rec = Reconciliation(
        activity_id=act.activity_id,
        reconciled_status=ActivityStatus.IN_PROGRESS,
        reconciled_progress=40.0,
        actual_start=date(2026, 9, 5),
        actual_end=None,
        confidence=0.72,
        explanation="Moderate confidence evidence requires planner verification.",
        decision=ReconciliationDecision.PLANNER_REVIEW,
        supporting_evidence=[{"source_name": "DPR-Unit3.txt"}],
    )
    session.add(rec)
    session.flush()

    review = PlannerReview(
        activity_id=act.activity_id,
        reconciliation_id=rec.reconciliation_id,
        reason=rec.explanation,
        proposed_decision=ReviewDecision.ACCEPT,
        reviewer_decision=ReviewDecision.PENDING,
    )
    session.add(review)
    session.commit()

    service = ScheduleBridgeService(session)
    preview = service.get_change_preview(bridge_project.project_id)

    assert preview.summary.planner_review == 1
    assert preview.summary.verified_updates == 0

    csv_bytes, _, _ = service.export_verified_schedule(bridge_project.project_id, "csv")
    csv_text = csv_bytes.decode("utf-8-sig")
    reader = list(csv.DictReader(io.StringIO(csv_text)))

    assert len(reader) == 1
    row = reader[0]
    assert row["activity_id"] == "PIP002"
    assert row["export_status"] == "EXCLUDED_PLANNER_REVIEW"
    # CRITICAL INVARIANT: Proposed unverified actual_start MUST NOT appear in export!
    assert row["actual_start"] == ""
    assert row["actual_end"] == ""
    assert row["percent_complete"] == ""


# ---------------------------------------------------------------------------
# 6. Conflict Exclusion from Verified Export
# ---------------------------------------------------------------------------

def test_conflict_exclusion_from_verified_export(session, bridge_project):
    act = Activity(
        project_id=bridge_project.project_id,
        external_activity_id="ELEC002",
        description="Cable tray installation",
        discipline="Electrical",
        level=6,
        status=ActivityStatus.NOT_STARTED,
    )
    session.add(act)
    session.flush()

    conflict = EvidenceConflict(
        activity_id=act.activity_id,
        conflict_type=ConflictType.STATUS,
        severity=ConflictSeverity.HIGH,
        description="Daily report contradicts subcontractor return.",
    )
    session.add(conflict)
    session.commit()

    service = ScheduleBridgeService(session)
    preview = service.get_change_preview(bridge_project.project_id)

    assert preview.summary.conflicts == 1
    assert preview.summary.verified_updates == 0

    csv_bytes, _, _ = service.export_verified_schedule(bridge_project.project_id, "csv")
    csv_text = csv_bytes.decode("utf-8-sig")
    row = list(csv.DictReader(io.StringIO(csv_text)))[0]

    assert row["export_status"] == "EXCLUDED_CONFLICT"
    assert row["actual_start"] == ""
    assert row["actual_end"] == ""


# ---------------------------------------------------------------------------
# 7. Rejected Exclusion from Verified Export
# ---------------------------------------------------------------------------

def test_rejected_exclusion_from_verified_export(session, bridge_project):
    act = Activity(
        project_id=bridge_project.project_id,
        external_activity_id="CIV003",
        description="Fencing outside battery limit",
        level=6,
        status=ActivityStatus.NOT_STARTED,
    )
    session.add(act)
    session.flush()

    rec = Reconciliation(
        activity_id=act.activity_id,
        reconciled_status=ActivityStatus.UNKNOWN,
        confidence=0.1,
        explanation="Uninterpretable evidence.",
        decision=ReconciliationDecision.REJECT,
    )
    session.add(rec)
    session.commit()

    service = ScheduleBridgeService(session)
    preview = service.get_change_preview(bridge_project.project_id)

    assert preview.summary.rejected == 1

    csv_bytes, _, _ = service.export_verified_schedule(bridge_project.project_id, "csv")
    csv_text = csv_bytes.decode("utf-8-sig")
    row = list(csv.DictReader(io.StringIO(csv_text)))[0]

    assert row["export_status"] == "EXCLUDED_REJECTED"
    assert row["actual_start"] == ""
    assert row["actual_end"] == ""


# ---------------------------------------------------------------------------
# 8. Unchanged Activity Handling
# ---------------------------------------------------------------------------

def test_unchanged_activity_handling(session, bridge_project):
    act = Activity(
        project_id=bridge_project.project_id,
        external_activity_id="INST001",
        description="Instrument junction box installation",
        discipline="Instrumentation",
        level=6,
        status=ActivityStatus.NOT_STARTED,
        planned_start=date(2026, 9, 1),
        planned_finish=date(2026, 9, 15),
    )
    session.add(act)
    session.commit()

    service = ScheduleBridgeService(session)
    preview = service.get_change_preview(bridge_project.project_id)

    assert preview.summary.unchanged == 1
    assert preview.summary.verified_updates == 0

    csv_bytes, _, _ = service.export_verified_schedule(bridge_project.project_id, "csv")
    csv_text = csv_bytes.decode("utf-8-sig")
    row = list(csv.DictReader(io.StringIO(csv_text)))[0]

    assert row["export_status"] == "UNCHANGED"
    assert row["actual_start"] == ""
    assert row["actual_end"] == ""
    assert row["status"] == "not_started"


# ---------------------------------------------------------------------------
# 9. Export Field Correctness (CSV)
# ---------------------------------------------------------------------------

def test_export_field_correctness_csv(session, bridge_project):
    act = Activity(
        project_id=bridge_project.project_id,
        external_activity_id="PIP001",
        description="Erect Line 24-XX",
        discipline="Piping",
        location="Unit 3",
        level=6,
        status=ActivityStatus.NOT_STARTED,
        planned_start=date(2026, 8, 25),
        planned_finish=date(2026, 9, 20),
    )
    session.add(act)
    session.flush()

    rec = Reconciliation(
        activity_id=act.activity_id,
        reconciled_status=ActivityStatus.COMPLETED,
        reconciled_progress=100.0,
        actual_start=date(2026, 8, 25),
        actual_end=date(2026, 9, 12),
        confidence=0.98,
        explanation="100% erection completed.",
        decision=ReconciliationDecision.AUTO_ACCEPT,
        supporting_evidence=[{"source_name": "DPR-104.txt"}],
    )
    session.add(rec)
    session.flush()

    ScheduleUpdateService(session).apply(
        act, rec, decision_source="automatic:reconciliation", reason=rec.explanation
    )
    session.commit()

    service = ScheduleBridgeService(session)
    csv_bytes, media_type, filename = service.export_verified_schedule(bridge_project.project_id, "csv")

    assert media_type.startswith("text/csv")
    assert filename.endswith(".csv")

    csv_text = csv_bytes.decode("utf-8-sig")
    reader = list(csv.DictReader(io.StringIO(csv_text)))
    assert len(reader) == 1
    row = reader[0]

    assert row["activity_id"] == "PIP001"
    assert row["activity_name"] == "Erect Line 24-XX"
    assert row["discipline"] == "Piping"
    assert row["location"] == "Unit 3"
    assert row["planned_start"] == "2026-08-25"
    assert row["planned_finish"] == "2026-09-20"
    assert row["actual_start"] == "2026-08-25"
    assert row["actual_end"] == "2026-09-12"
    assert row["percent_complete"] == "100.0"
    assert row["status"] == "completed"
    assert row["export_status"] == "VERIFIED_UPDATE"
    assert row["verification_source"] == "automatic:reconciliation"
    assert row["source_evidence"] == "DPR-104.txt"


# ---------------------------------------------------------------------------
# 10. Export Field Correctness (XLSX)
# ---------------------------------------------------------------------------

def test_export_field_correctness_xlsx(session, bridge_project):
    act = Activity(
        project_id=bridge_project.project_id,
        external_activity_id="PIP001",
        description="Erect Line 24-XX",
        discipline="Piping",
        level=6,
        status=ActivityStatus.NOT_STARTED,
    )
    session.add(act)
    session.flush()

    rec = Reconciliation(
        activity_id=act.activity_id,
        reconciled_status=ActivityStatus.COMPLETED,
        reconciled_progress=100.0,
        actual_start=date(2026, 8, 25),
        actual_end=date(2026, 9, 12),
        confidence=0.95,
        explanation="Completed erection.",
        decision=ReconciliationDecision.AUTO_ACCEPT,
        supporting_evidence=[{"source_name": "DPR-104.txt"}],
    )
    session.add(rec)
    session.flush()

    ScheduleUpdateService(session).apply(
        act, rec, decision_source="automatic:reconciliation", reason=rec.explanation
    )
    session.commit()

    service = ScheduleBridgeService(session)
    xlsx_bytes, media_type, filename = service.export_verified_schedule(bridge_project.project_id, "xlsx")

    assert "spreadsheetml" in media_type
    assert filename.endswith(".xlsx")

    wb = load_workbook(io.BytesIO(xlsx_bytes))
    assert "Verified Schedule" in wb.sheetnames
    assert "Review & Conflicts Audit" in wb.sheetnames

    ws = wb["Verified Schedule"]
    assert ws.max_row == 2  # Header + 1 activity
    assert ws.cell(row=2, column=1).value == "PIP001"
    assert ws.cell(row=2, column=9).value == "2026-08-25"
    assert ws.cell(row=2, column=10).value == "2026-09-12"
    assert ws.cell(row=2, column=13).value == "VERIFIED_UPDATE"


# ---------------------------------------------------------------------------
# 11. Export Audit and Provenance Traceability
# ---------------------------------------------------------------------------

def test_export_audit_provenance(session, bridge_project):
    act = Activity(
        project_id=bridge_project.project_id,
        external_activity_id="PIP001",
        description="Erect Line 24-XX",
        level=6,
        status=ActivityStatus.NOT_STARTED,
    )
    session.add(act)
    session.flush()

    evidence = Evidence(
        project_id=bridge_project.project_id,
        source_type=SourceType.DAILY_REPORT,
        source_name="DPR-Unit3-0912.txt",
        source_timestamp=datetime(2026, 9, 12, 16, 0, tzinfo=UTC),
        raw_text="Piping crew completed Line 24-XX at 16:00.",
    )
    session.add(evidence)
    session.flush()

    obs = EvidenceInterpretation(
        evidence_id=evidence.evidence_id,
        original_evidence=evidence.raw_text,
        activity_description="Line 24-XX erection",
        status=ActivityStatus.COMPLETED,
        actual_end=date(2026, 9, 12),
        confidence=0.98,
        provider_name="test_provider",
        raw_provider_response={},
    )
    session.add(obs)
    session.flush()

    rec = Reconciliation(
        activity_id=act.activity_id,
        reconciled_status=ActivityStatus.COMPLETED,
        reconciled_progress=100.0,
        actual_end=date(2026, 9, 12),
        confidence=0.98,
        explanation="Field DPR verifies completion.",
        decision=ReconciliationDecision.AUTO_ACCEPT,
        supporting_evidence=[{
            "observation_id": str(obs.observation_id),
            "evidence_id": str(evidence.evidence_id),
            "source_name": evidence.source_name,
            "source_type": evidence.source_type.value,
        }],
    )
    session.add(rec)
    session.flush()

    update = ScheduleUpdateService(session).apply(
        act, rec, decision_source="automatic:reconciliation", reason=rec.explanation
    )
    session.commit()

    # Traceability check
    persisted_update = session.get(ScheduleUpdate, update.update_id)
    assert persisted_update.reconciliation_id == rec.reconciliation_id
    assert len(persisted_update.evidence) == 1
    assert persisted_update.evidence[0]["source_name"] == "DPR-Unit3-0912.txt"
    assert persisted_update.evidence[0]["evidence_id"] == str(evidence.evidence_id)


# ---------------------------------------------------------------------------
# 12. End-to-End Verified Schedule Flow (Scenario 1)
# ---------------------------------------------------------------------------

def test_e2e_verified_schedule_flow(session, bridge_project, tmp_path):
    sched_csv = tmp_path / "sched.csv"
    sched_csv.write_text(
        "activity_id,activity_name,discipline,planned_start,planned_end\n"
        "PIP001,Erect Line 24-XX,Piping,2026-09-12,2026-09-14\n",
        encoding="utf-8",
    )
    ingestion = IngestionService(session)
    ingestion.import_schedule(bridge_project.project_id, sched_csv, "sched.csv", "v1")

    # Ingest text DPR
    dpr_text = "Piping crew completed erection of Line 24-XX at 16:00 on 12 Sep."
    _, evidence = ingestion.import_text_evidence(
        bridge_project.project_id, dpr_text, "dpr_sep12.txt", SourceType.DAILY_REPORT
    )

    act = session.scalar(
        select(Activity).where(
            Activity.project_id == bridge_project.project_id,
            Activity.external_activity_id == "PIP001",
        )
    )
    assert act is not None

    obs = EvidenceInterpretation(
        evidence_id=evidence.evidence_id,
        original_evidence=dpr_text,
        activity_description="Erect Line 24-XX",
        status=ActivityStatus.COMPLETED,
        actual_end=date(2026, 9, 12),
        confidence=0.98,
        provider_name="test_provider",
        raw_provider_response={},
    )
    session.add(obs)
    session.flush()

    # Match observation to PIP001
    result = ObservationMatchResult(
        observation_id=obs.observation_id,
        outcome=ObservationMatchOutcome.MATCHED,
        best_score=0.95,
        high_confidence_threshold=0.80,
        ambiguous_score_delta=0.08,
        no_match_threshold=0.50,
    )
    session.add(result)
    session.flush()

    candidate = ObservationMatchCandidate(
        result_id=result.result_id,
        activity_id=act.activity_id,
        rank=1,
        similarity_score=0.95,
        contextual_score=0.95,
        final_score=0.95,
        explanation="Exact match on Line 24-XX",
    )
    session.add(candidate)
    session.commit()

    # Reconcile activity
    class MockE2EProvider:
        provider_name = "mock_e2e_provider"
        model_name = "mock-v1"

        def explain_reconciliation(self, prompt: str) -> str:
            return json.dumps({
                "recommended_status": "completed",
                "recommended_progress": 100.0,
                "recommended_actual_end": "2026-09-12",
                "confidence": 0.98,
                "explanation": "Piping crew completed erection on Sep 12.",
                "conflict_detected": False,
            })

    rec_service = ReconciliationService(session, llm_provider=MockE2EProvider())
    rec_response = rec_service.reconcile_activity(act.activity_id)

    assert rec_response.decision == ReconciliationDecision.AUTO_ACCEPT
    assert act.status == ActivityStatus.COMPLETED
    assert act.actual_end == date(2026, 9, 12)

    # Export verified schedule
    bridge_service = ScheduleBridgeService(session)
    preview = bridge_service.get_change_preview(bridge_project.project_id)
    assert preview.summary.verified_updates == 1

    csv_bytes, _, _ = bridge_service.export_verified_schedule(bridge_project.project_id, "csv")
    csv_text = csv_bytes.decode("utf-8-sig")
    row = list(csv.DictReader(io.StringIO(csv_text)))[0]

    assert row["activity_id"] == "PIP001"
    assert row["actual_end"] == "2026-09-12"
    assert row["export_status"] == "VERIFIED_UPDATE"


# ---------------------------------------------------------------------------
# 13. Dependency Conflict Cannot Enter Verified Export (Scenario 2)
# ---------------------------------------------------------------------------

def test_dependency_conflict_cannot_enter_verified_export(session, bridge_project, tmp_path):
    sched_csv = tmp_path / "dep_sched.csv"
    sched_csv.write_text(
        "activity_id,activity_name,discipline,planned_start,planned_end,predecessors\n"
        "PIP001,Erect Line 24-XX,Piping,2026-08-25,2026-09-20,\n"
        "HT001,Hydrotest Line 24-XX,Piping,2026-09-21,2026-09-25,PIP001\n",
        encoding="utf-8",
    )
    ingestion = IngestionService(session)
    ingestion.import_schedule(bridge_project.project_id, sched_csv, "sched.csv", "v1")

    # Ingest DPR claiming HT001 commenced on 2026-09-10 (violates PIP001 completion on 2026-09-15!)
    pip = session.scalar(
        select(Activity).where(
            Activity.project_id == bridge_project.project_id,
            Activity.external_activity_id == "PIP001",
        )
    )
    assert pip is not None
    pip.status = ActivityStatus.COMPLETED
    pip.actual_start = date(2026, 9, 1)
    pip.actual_end = date(2026, 9, 15)
    pip.actual_progress = 100.0
    session.flush()

    _, evidence = ingestion.import_text_evidence(
        bridge_project.project_id,
        "Hydrotest crew commenced testing on Sep 10.",
        "dpr_ht.txt",
        SourceType.DAILY_REPORT,
    )

    ht = session.scalar(
        select(Activity).where(
            Activity.project_id == bridge_project.project_id,
            Activity.external_activity_id == "HT001",
        )
    )

    obs = EvidenceInterpretation(
        evidence_id=evidence.evidence_id,
        original_evidence="Hydrotest started Sep 10",
        activity_description="Hydrotest Line 24-XX",
        status=ActivityStatus.IN_PROGRESS,
        actual_start=date(2026, 9, 10),
        confidence=0.98,
        provider_name="test_provider",
        raw_provider_response={},
    )
    session.add(obs)
    session.flush()

    res = ObservationMatchResult(
        observation_id=obs.observation_id,
        outcome=ObservationMatchOutcome.MATCHED,
        best_score=0.95,
        high_confidence_threshold=0.8,
        ambiguous_score_delta=0.08,
        no_match_threshold=0.5,
    )
    session.add(res)
    session.flush()
    session.add(
        ObservationMatchCandidate(
            result_id=res.result_id,
            activity_id=ht.activity_id,
            rank=1,
            similarity_score=0.95,
            contextual_score=0.95,
            final_score=0.95,
            explanation="Hydrotest match",
        )
    )
    session.commit()

    # Reconcile HT001: safety gate MUST block AUTO_ACCEPT due to dependency conflict
    from app.ai.mock_provider import DeterministicMockProvider
    rec_service = ReconciliationService(session, llm_provider=DeterministicMockProvider())
    rec_response = rec_service.reconcile_activity(ht.activity_id)

    assert rec_response.decision == ReconciliationDecision.PLANNER_REVIEW

    # Export verified schedule: HT001 actuals MUST NOT be exported!
    bridge_service = ScheduleBridgeService(session)
    preview = bridge_service.get_change_preview(bridge_project.project_id)
    assert preview.summary.verified_updates == 0
    assert preview.summary.conflicts >= 1 or preview.summary.planner_review >= 1

    csv_bytes, _, _ = bridge_service.export_verified_schedule(bridge_project.project_id, "csv")
    csv_text = csv_bytes.decode("utf-8-sig")
    ht_row = next(r for r in csv.DictReader(io.StringIO(csv_text)) if r["activity_id"] == "HT001")

    assert "EXCLUDED" in ht_row["export_status"]
    assert ht_row["actual_start"] == ""
    assert ht_row["actual_end"] == ""


# ---------------------------------------------------------------------------
# 14. Ambiguous Match Cannot Enter Verified Export (Scenario 3)
# ---------------------------------------------------------------------------

def test_ambiguous_match_cannot_enter_verified_export(session, bridge_project, tmp_path):
    sched_csv = tmp_path / "ambig_sched.csv"
    sched_csv.write_text(
        "activity_id,activity_name,discipline,location,planned_start,planned_end\n"
        "ELEC001,Cable tray installation,Electrical,Unit 1,2026-09-01,2026-09-18\n"
        "ELEC002,Cable tray installation,Electrical,Unit 2,2026-09-01,2026-09-18\n",
        encoding="utf-8",
    )
    ingestion = IngestionService(session)
    ingestion.import_schedule(bridge_project.project_id, sched_csv, "sched.csv", "v1")

    # Ingest observation with ambiguous match
    _, evidence = ingestion.import_text_evidence(
        bridge_project.project_id,
        "Completed cable tray installation.",
        "dpr_elec.txt",
        SourceType.DAILY_REPORT,
    )

    obs = EvidenceInterpretation(
        evidence_id=evidence.evidence_id,
        original_evidence="Completed cable tray installation.",
        activity_description="Cable tray installation",
        status=ActivityStatus.COMPLETED,
        actual_end=date(2026, 9, 15),
        confidence=0.95,
        provider_name="test_provider",
        raw_provider_response={},
    )
    session.add(obs)
    session.flush()

    # Ambiguous match outcome
    res = ObservationMatchResult(
        observation_id=obs.observation_id,
        outcome=ObservationMatchOutcome.AMBIGUOUS,
        ambiguity_flag=True,
        best_score=0.88,
        high_confidence_threshold=0.8,
        ambiguous_score_delta=0.08,
        no_match_threshold=0.5,
    )
    session.add(res)
    session.commit()

    # Neither ELEC001 nor ELEC002 has an auto-accepted update
    bridge_service = ScheduleBridgeService(session)
    preview = bridge_service.get_change_preview(bridge_project.project_id)

    assert preview.summary.verified_updates == 0

    csv_bytes, _, _ = bridge_service.export_verified_schedule(bridge_project.project_id, "csv")
    csv_text = csv_bytes.decode("utf-8-sig")
    rows = list(csv.DictReader(io.StringIO(csv_text)))

    for row in rows:
        assert row["actual_end"] == ""
        assert row["export_status"] == "UNCHANGED"


# ---------------------------------------------------------------------------
# 15. API Endpoints Test
# ---------------------------------------------------------------------------

def test_schedule_api_preview_and_export(session, bridge_project):
    from app.api.routes.imports import get_session

    app.dependency_overrides[get_session] = lambda: session
    try:
        client = TestClient(app)

        # Preview endpoint
        preview_res = client.get(f"/api/v1/projects/{bridge_project.project_id}/schedule/preview")
        assert preview_res.status_code == 200
        data = preview_res.json()
        assert "summary" in data
        assert "changes" in data
        assert "P6 XML deferred" in data["deferral_notice"]

        # CSV Export endpoint
        csv_res = client.get(f"/api/v1/projects/{bridge_project.project_id}/schedule/export?format=csv")
        assert csv_res.status_code == 200
        assert csv_res.headers["content-type"].startswith("text/csv")
        assert "attachment; filename=" in csv_res.headers["content-disposition"]

        # XLSX Export endpoint
        xlsx_res = client.get(f"/api/v1/projects/{bridge_project.project_id}/schedule/export?format=xlsx")
        assert xlsx_res.status_code == 200
        assert "spreadsheetml" in xlsx_res.headers["content-type"]
        assert "attachment; filename=" in xlsx_res.headers["content-disposition"]
    finally:
        app.dependency_overrides.clear()
