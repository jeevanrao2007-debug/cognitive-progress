import csv
from collections.abc import Iterable
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from app.modules.ingestion.errors import IngestionError


def read_tabular_rows(path: Path) -> list[dict[str, str]]:
    """Read CSV or XLSX headers and values without interpreting the domain fields."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            return [
                {str(key).strip(): "" if value is None else str(value).strip() for key, value in row.items() if key}
                for row in csv.DictReader(source)
            ]
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook.active
        iterator: Iterable[tuple[object, ...]] = worksheet.iter_rows(values_only=True)
        try:
            headers = next(iterator)
        except StopIteration as error:
            raise IngestionError("Spreadsheet is empty.") from error
        normalized_headers = [str(header).strip() if header is not None else "" for header in headers]
        if not any(normalized_headers):
            raise IngestionError("Spreadsheet has no header row.")
        return [
            {
                header: "" if value is None else str(value).strip()
                for header, value in zip(normalized_headers, row, strict=False)
                if header
            }
            for row in iterator
            if any(value is not None and str(value).strip() for value in row)
        ]
    raise IngestionError("Only .csv and .xlsx tabular uploads are supported.")


def read_text(path: Path) -> str:
    """Read text directly or extract selectable text from a PDF; OCR is out of scope."""
    suffix = path.suffix.lower()
    if suffix == ".txt":
        try:
            return path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            try:
                return path.read_text(encoding="utf-8-sig")
            except UnicodeDecodeError:
                return path.read_text(encoding="latin-1")
    if suffix == ".pdf":
        try:
            extracted = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages).strip()
            if not extracted:
                raise IngestionError("The uploaded PDF document contains no selectable text.")
            return extracted
        except IngestionError:
            raise
        except Exception as error:  # pypdf exposes several parser-specific exception classes.
            raise IngestionError("Could not extract selectable text from the PDF.") from error
    raise IngestionError("Only .txt and text-based .pdf uploads are supported.")

